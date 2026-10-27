"""Build a journal-formatted .docx from draft.md (format described in SKILL.md, step 6).

Usage:
  python -I build_docx.py draft.md out.docx [--profile mcu|general] [--font "TH Sarabun PSK"]
                          [--body-size 16] [--heading-size 18] [--title-size 20] [--indent-cm 1.27]
                          [--align thai|justify|left]

mcu profile follows the common layout of MCU-affiliated Thai journals: A4, 1-inch margins, single spacing, no page numbers,
type label 12pt, Thai title 20pt, English title 18pt caps, authors 14pt, headings 18pt, body 16pt,
table text 14pt, Thai-distributed body paragraphs, hanging-indent references.
Image paths in ![caption](path) are resolved relative to draft.md.
"""
import argparse, os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

REF_HEADS = {"เอกสารอ้างอิง", "บรรณานุกรม", "References", "Reference"}
THAI = re.compile(r"[฀-๿]")
# Compounds the default pythainlp dictionary splits, which would let Word wrap mid-word (e.g. ตัวชี้ / วัด).
KEEP_WORDS = {"ตัวชี้วัด", "ผู้ทำบัญชี", "ผู้ทรงคุณวุฒิ", "ผู้ให้ข้อมูล", "ผู้ให้บริการ", "ผู้เสียภาษี",
              "ใบกำกับภาษี", "ภาษีมูลค่าเพิ่ม", "ภาษีเงินได้", "นิติบุคคล", "บทคัดย่อ", "คำสำคัญ",
              "องค์ความรู้", "หลักพุทธธรรม", "พระไตรปิฎก", "สถิติเชิงพรรณนา", "สัมประสิทธิ์"}
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


class Builder:
    def __init__(self, a):
        self.a = a
        self.tokenize = None
        if a.thai_wordbreak != "off":
            try:
                from pythainlp.tokenize import word_tokenize
                from pythainlp.corpus.common import thai_words
                from pythainlp.util import Trie
                words = set(thai_words()) | KEEP_WORDS
                if a.thai_dict:
                    with open(a.thai_dict, encoding="utf-8") as f:
                        words |= {w.strip() for w in f if w.strip()}
                trie = Trie(words)
                self.tokenize = lambda t, **kw: word_tokenize(t, custom_dict=trie, **kw)
            except ImportError:
                if a.thai_wordbreak == "on":
                    sys.exit("--thai-wordbreak on needs: pip install pythainlp")
                print("NOTE pythainlp not installed: Thai lines may wrap only at spaces in Word", file=sys.stderr)
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        st = self.doc.styles["Normal"]
        self._font(st.element.get_or_add_rPr(), a.body_size)
        pf = st.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # --- low-level formatting -------------------------------------------------
    def _font(self, rpr, size, bold=False, italic=False):
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(k), self.a.font)
        for tag in ("w:b", "w:bCs", "w:i", "w:iCs", "w:sz", "w:szCs", "w:lang"):
            el = rpr.find(qn(tag))
            if el is not None:
                rpr.remove(el)
        if bold:
            rpr.append(OxmlElement("w:b")); rpr.append(OxmlElement("w:bCs"))
        if italic:
            rpr.append(OxmlElement("w:i")); rpr.append(OxmlElement("w:iCs"))
        for tag in ("w:sz", "w:szCs"):
            el = OxmlElement(tag); el.set(qn("w:val"), str(int(size * 2))); rpr.append(el)
        lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "th-TH"); rpr.append(lang)

    def wordbreak(self, text):
        """Insert zero-width spaces between Thai words so Word can wrap lines without a Thai word-breaker."""
        if not self.tokenize or not THAI.search(text):
            return text
        out = []
        for t in self.tokenize(text, keep_whitespace=True):
            if not t:
                continue
            # no break opportunity right after an opening bracket, so "(" never ends a line
            if out and not out[-1].endswith(("(", "[")):
                out.append("​")
            out.append(t)
        return "".join(out)

    def run(self, p, text, size, bold=False, italic=False, sup=False):
        r = p.add_run(self.wordbreak(text))
        self._font(r._r.get_or_add_rPr(), size, bold, italic)
        if sup:
            r.font.superscript = True
        return r

    def para(self, align=None, indent=False, hanging=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        if align == "thai":
            jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "thaiDistribute")
            p._p.get_or_add_pPr().append(jc)
        elif align is not None:
            p.alignment = align
        if indent:
            pf.first_line_indent = Cm(self.a.indent_cm)
        if hanging:
            pf.left_indent = Cm(self.a.indent_cm)
            pf.first_line_indent = Cm(-self.a.indent_cm)
        return p

    def rich(self, p, text, size, bold=False):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**"):
                self.run(p, part[2:-2], size, bold=True)
            elif part.startswith("*") and len(part) > 2:
                self.run(p, part[1:-1], size, bold=bold, italic=True)
            else:
                self.run(p, part, size, bold=bold)

    def marked(self, p, text, size, bold=False):
        """Author/affiliation text: trailing digits and * become superscript."""
        for part in re.split(r"(\d+\*?|\*)", text):
            if not part:
                continue
            self.run(p, part, size, bold=bold, sup=bool(re.fullmatch(r"\d+\*?|\*", part)))

    # --- blocks ---------------------------------------------------------------
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def page_numbers(self):
        if self.a.page_numbers == "none":
            return
        hdr = self.doc.sections[0].header
        p = hdr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if self.a.page_numbers == "top-right" else WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        self._font(r._r.get_or_add_rPr(), self.a.body_size)
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
            r._r.append(el)

    def header_left(self, meta, lang):
        """Title centred; authors, affiliations and e-mail left-aligned (one language block)."""
        a, size = self.a, self.a.body_size
        title = meta.get(f"title_{lang}")
        if title:
            if lang == "en" and a.title_en_upper:
                title = title.upper()
            self.run(self.para(WD_ALIGN_PARAGRAPH.CENTER), title, a.title_en_size if lang == "en" else a.title_size, bold=True)
            self.para()
        if meta.get(f"authors_{lang}"):
            self.marked(self.para(WD_ALIGN_PARAGRAPH.LEFT), meta[f"authors_{lang}"], size, bold=True)
        for line in filter(None, (x.strip() for x in meta.get(f"affiliations_{lang}", "").split("|"))):
            self.marked(self.para(WD_ALIGN_PARAGRAPH.LEFT), line, size)
        if meta.get("email"):
            self.run(self.para(WD_ALIGN_PARAGRAPH.LEFT), f"Email: {meta['email']}", size)

    def header(self, meta):
        a, C = self.a, WD_ALIGN_PARAGRAPH.CENTER
        if a.header_style == "left":
            self.header_left(meta, "th")
            if not a.split_en_header:
                self.header_left(meta, "en")
            return
        if a.profile == "mcu" and meta.get("type"):
            self.run(self.para(WD_ALIGN_PARAGRAPH.LEFT), meta["type"], 12, bold=True)
        if meta.get("title_th"):
            self.run(self.para(C), meta["title_th"], a.title_size, bold=True)
        if meta.get("title_en"):
            self.run(self.para(C), meta["title_en"].upper() if a.title_en_upper else meta["title_en"], a.title_en_size, bold=True)
        for key in ("authors_th", "authors_en"):
            if meta.get(key):
                self.marked(self.para(C), meta[key], 14)
        for key in ("affiliations_th", "affiliations_en"):
            for line in filter(None, (x.strip() for x in meta.get(key, "").split("|"))):
                self.marked(self.para(C), line, 14)
        if meta.get("email"):
            p = self.para(C)
            self.run(p, "*", 14, sup=True)
            self.run(p, f"Corresponding Author E-mail: {meta['email']}", 14)
        if a.profile == "mcu":
            self.run(self.para(C), meta.get("received", "Received……; Revised……; Accepted……"), 14)

    def image(self, caption, path, base):
        full = os.path.join(base, path)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        if os.path.exists(full):
            p.add_run().add_picture(full, width=Inches(6))
            p.paragraph_format.keep_with_next = True
        else:
            self.run(p, f"[ไม่พบไฟล์ภาพ: {path}]", self.a.body_size)
            print(f"WARN missing image {full}", file=sys.stderr)
        cap = self.para(WD_ALIGN_PARAGRAPH.LEFT if self.a.caption_align == "left" else WD_ALIGN_PARAGRAPH.CENTER)
        m = re.match(r"(ภาพที่\s*\d+|Figure\s*\d+)(.*)", caption)
        if m:
            self.run(cap, m.group(1), self.a.body_size, bold=True)
            self.run(cap, m.group(2), self.a.body_size)
        else:
            self.run(cap, caption, self.a.body_size)

    def table(self, rows):
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
        if not cells:
            return
        ncol = max(len(r) for r in cells)
        t = self.doc.add_table(rows=len(cells), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if self.a.table_style == "horizontal":
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{edge}")
                on = edge in ("top", "bottom", "insideH")
                el.set(qn("w:val"), "single" if on else "nil")
                if on:
                    el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
                borders.append(el)
            tblPr = t._tbl.tblPr
            old = tblPr.find(qn("w:tblBorders"))
            if old is not None:
                tblPr.remove(old)
            tblPr.append(borders)
        for i, r in enumerate(cells):
            trPr = t.rows[i]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trPr.append(OxmlElement("w:tblHeader"))  # repeat header row on each page
            for j in range(ncol):
                p = t.cell(i, j).paragraphs[0]
                self.rich(p, r[j] if j < len(r) else "", 14, bold=(i == 0))
                # header stays with the first data row; last row stays with the source line
                if i == 0 or i == len(cells) - 1:
                    p.paragraph_format.keep_with_next = True
        self.para()

    def build(self, md_path, out_path):
        text = open(md_path, encoding="utf-8").read()
        meta = {}
        if text.startswith("---"):
            end = text.find("\n---", 3)
            for line in text[3:end].splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    meta[k.strip()] = v.strip()
            text = text[end + 4:]
        text = re.sub(r"\[/?ตัวบท\]", "", text)  # verbatim-law markers used by thai_style_check.py
        self.page_numbers()
        self.header(meta)
        breaks = [x.strip() for x in self.a.page_break_before.split("|") if x.strip()]
        centered = [x.strip() for x in self.a.center_headings.split("|") if x.strip()]
        base = os.path.dirname(os.path.abspath(md_path))
        lines, i, in_refs, size = text.splitlines(), 0, False, self.a.body_size
        while i < len(lines):
            line = lines[i].rstrip()
            s = line.strip()
            i += 1
            if not s:
                continue
            if s.startswith("# "):
                title = s[2:].strip()
                in_refs = title in REF_HEADS or title.startswith("เอกสารอ้างอิง") or title.startswith("References")
                if self.a.split_en_header and title == "Abstract":
                    self.page_break()
                    self.header_left(meta, "en")
                    self.para()
                new_page = not (self.a.split_en_header and title == "Abstract") and any(title.startswith(b) for b in breaks)
                if not new_page and not (self.a.split_en_header and title == "Abstract"):
                    self.para()
                align = WD_ALIGN_PARAGRAPH.CENTER if title in centered else WD_ALIGN_PARAGRAPH.LEFT
                hp = self.para(align)
                hp.paragraph_format.keep_with_next = True
                # page_break_before on the heading itself: a separate break paragraph can spill onto
                # its own page when the previous page is exactly full, leaving a blank page
                hp.paragraph_format.page_break_before = new_page
                self.run(hp, title, self.a.heading_size, bold=True)
                continue
            if s.startswith("## ") or s.startswith("### "):
                hp = self.para(WD_ALIGN_PARAGRAPH.LEFT)
                hp.paragraph_format.keep_with_next = True
                self.rich(hp, s.lstrip("#").strip(), size, bold=True)
                continue
            m = re.fullmatch(r"!\[(.*?)\]\((.+?)\)", s)
            if m:
                self.image(m.group(1), m.group(2), base)
                continue
            if s.startswith("|"):
                rows = [s]
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(lines[i].strip()); i += 1
                self.table(rows)
                continue
            # merge soft-wrapped lines into one paragraph
            buf = [s]
            while i < len(lines) and lines[i].strip() and not re.match(r"(#|!\[|\||ที่มา:|ตารางที่|คำสำคัญ:|Keywords:)", lines[i].strip()):
                buf.append(lines[i].strip()); i += 1
            s = " ".join(buf)
            km = re.match(r"^(คำสำคัญ|Keywords)\s*:\s*(.*)$", s)
            if km:
                p = self.para(WD_ALIGN_PARAGRAPH.LEFT)
                self.run(p, km.group(1) + ": ", size, bold=True)
                self.run(p, km.group(2), size)
            elif s.startswith("ที่มา:") or s.startswith("Source:"):
                self.rich(self.para(WD_ALIGN_PARAGRAPH.LEFT if self.a.caption_align == "left" else WD_ALIGN_PARAGRAPH.CENTER), s, size)
            elif re.match(r"^(ตารางที่|Table)\s*\d+", s):
                p = self.para(WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.keep_with_next = True
                m2 = re.match(r"^((?:ตารางที่|Table)\s*\d+)(.*)$", s)
                self.run(p, m2.group(1), size, bold=True)
                self.run(p, m2.group(2), size)
            elif in_refs:
                self.rich(self.para(WD_ALIGN_PARAGRAPH.LEFT, hanging=True), re.sub(r"^[-*]\s+", "", s), size)
            else:
                align = {"thai": "thai", "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}[self.a.align]
                self.rich(self.para(align, indent=True), s, size)
        self.doc.save(out_path)
        print(f"saved {out_path}")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser()
    p.add_argument("draft"); p.add_argument("out")
    p.add_argument("--profile", choices=["mcu", "general"], default="mcu")
    p.add_argument("--font", default="TH Sarabun PSK")
    p.add_argument("--body-size", type=float, default=16)
    p.add_argument("--heading-size", type=float, default=18)
    p.add_argument("--title-size", type=float, default=20)
    p.add_argument("--indent-cm", type=float, default=1.27)
    p.add_argument("--align", choices=["thai", "justify", "left"], default="thai",
                   help="body alignment: thai = Word 'Thai Distributed' (default), justify, left")
    p.add_argument("--title-en-size", type=float, help="English title size (default: title-size - 2)")
    p.add_argument("--title-en-upper", choices=["yes", "no"], default="yes")
    p.add_argument("--header-style", choices=["center", "left"], default="center",
                   help="center = MCU style; left = title centred, authors/affiliation/email left-aligned")
    p.add_argument("--split-en-header", action="store_true",
                   help="put the English title/authors on a new page right before '# Abstract'")
    p.add_argument("--page-numbers", choices=["none", "top-right", "bottom-center"], default="none")
    p.add_argument("--page-break-before", default="", help="'|'-separated heading prefixes that start on a new page")
    p.add_argument("--center-headings", default="", help="'|'-separated exact heading titles to centre")
    p.add_argument("--caption-align", choices=["center", "left"], default="center")
    p.add_argument("--table-style", choices=["grid", "horizontal"], default="grid",
                   help="horizontal = top/bottom/row rules only, no column lines")
    p.add_argument("--thai-wordbreak", choices=["auto", "on", "off"], default="auto",
                   help="insert zero-width spaces between Thai words (needs pythainlp) so Word wraps Thai text cleanly")
    p.add_argument("--thai-dict", help="UTF-8 file, one Thai word per line, kept unbroken when wrapping (adds to the built-in list)")
    a = p.parse_args()
    a.title_en_upper = a.title_en_upper == "yes"
    if a.title_en_size is None:
        a.title_en_size = a.title_size - 2
    Builder(a).build(a.draft, a.out)


if __name__ == "__main__":
    main()
